import io
from pptx import Presentation
from pptx.util import Inches

def generate_pptx(title: str, report_text: str, charts_images: list[bytes], insights: list[dict]) -> bytes:
    prs = Presentation()
    
    # Slide 1: Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title_shape = slide.shapes.title
    subtitle = slide.placeholders[1]
    title_shape.text = f"ManzStudio: {title}"
    subtitle.text = "Generated Report"
    
    # Slide 2: Executive Summary
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Executive Summary"
    tf = body_shape.text_frame
    tf.text = report_text[:500] + "..." if len(report_text) > 500 else report_text
    
    # Slides 3+: Charts
    blank_slide_layout = prs.slide_layouts[5]
    for i, img_bytes in enumerate(charts_images):
        slide = prs.slides.add_slide(blank_slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = f"Chart {i+1}"
        
        try:
            image_stream = io.BytesIO(img_bytes)
            slide.shapes.add_picture(image_stream, Inches(1), Inches(2), width=Inches(8))
        except:
            pass
            
        if i < len(insights):
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = insights[i].get('description', '')
            
    # Last slide: Key Recommendations
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Key Recommendations"
    tf = body_shape.text_frame
    recs = [ins for ins in insights if ins.get('type') == 'recommendation']
    for rec in recs:
        p = tf.add_paragraph()
        p.text = rec.get('title', '')
        
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
